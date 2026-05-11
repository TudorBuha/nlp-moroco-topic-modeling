"""Build the 10–15 minute project deck (PowerPoint)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MOROCO_topic_modeling_presentation.pptx"


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_bullets(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)


def _add_two_column(prs: Presentation, title: str, left: list[str], right: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5.0))
    right_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.5), Inches(4.5), Inches(5.0))
    for box, lines in ((left_box, left), (right_box, right)):
        tf = box.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(18)


def main() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(
        prs,
        "Topic Modeling on MOROCO",
        "BERTopic vs. LDA on Romanian news\nBuha Tudor · Ciorăscu Mihai",
    )

    _add_bullets(
        prs,
        "Problem & goal",
        [
            "Unsupervised topic discovery on Romanian news articles.",
            "Discover coherent word groups and document clusters without using gold labels at training time.",
            "Compare two approaches: classical LDA (Method A) and transformer-based BERTopic (Method B).",
            "Deliverables: trained models, evaluation, and an interactive Streamlit application.",
        ],
    )

    _add_bullets(
        prs,
        "Dataset — MOROCO",
        [
            "Moldavian and Romanian Dialectal Corpus (~33k news samples).",
            "Six topic categories: culture, finance, politics, science, sports, tech.",
            "Two dialects: Romanian and Moldavian (diacritics and spelling variants).",
            "Demo run: stratified 6k subset (1k per class), 80/20 train/test, seed 42.",
            "Same split for both methods so metrics are directly comparable.",
        ],
    )

    _add_two_column(
        prs,
        "Method A — LDA",
        [
            "Generative probabilistic model (Gensim).",
            "Input: bag-of-words after cleaning, stop-word removal, lemmatization.",
            "K chosen by sweeping C_v coherence (best K = 15).",
            "Every document gets a topic mixture.",
            "Strength: fast on CPU, familiar baseline.",
        ],
        [
            "Method B — BERTopic",
            "RoBERT embeddings → UMAP → HDBSCAN → c-TF-IDF.",
            "Topic count discovered from data (19 + outlier cluster).",
            "Romanian-specific encoder: readerbench/robert-base.",
            "Explicit outlier topic (-1) for ambiguous documents.",
            "Strength: semantics and dialect-sensitive clusters.",
        ],
    )

    _add_bullets(
        prs,
        "Evaluation",
        [
            "Intrinsic: C_v topic coherence on training topics.",
            "Extrinsic vs. MOROCO labels: NMI and Purity on the held-out test set.",
            "Confusion matrices: predicted topic vs. gold category.",
            "BERTopic extras: hyperparameter sweep, embedding ablation, multi-seed stability, bootstrap CIs.",
        ],
    )

    _add_bullets(
        prs,
        "Headline results (test set)",
        [
            "LDA (K=15): NMI 0.346, Purity 0.591, C_v 0.491.",
            "BERTopic (default): NMI 0.334, Purity 0.598, 19 topics, ~3.7% outliers on test.",
            "No single winner: LDA slightly higher NMI; BERTopic slightly higher Purity.",
            "BERTopic finds finer-grained themes (e.g. tech/privacy, weather, currency).",
            "Romanian RoBERT beats multilingual MPNet in ablation (more topics, higher Purity).",
        ],
    )

    _add_bullets(
        prs,
        "What we learned",
        [
            "LDA keywords are interpretable after lemmatization; topics align with broad MOROCO classes.",
            "BERTopic separates some themes LDA merges (e.g. Moldavian spelling in science clusters).",
            "Both models struggle with generic function words in Romanian if vectorization is too loose.",
            "UMAP/HDBSCAN topic count varies across seeds — report stability, not a single lucky run.",
            "Interactive labels in the GUI make exploration easier than raw topic ids alone.",
        ],
    )

    _add_bullets(
        prs,
        "Streamlit application",
        [
            "Home — headline metrics and pipeline status.",
            "Try it live — paste Romanian text; both models predict in parallel.",
            "LDA / BERTopic explorers — topic tables, coherence curve, pyLDAvis, Plotly HTML maps.",
            "Comparison — side-by-side metrics, confusion matrices, embedding ablation.",
            "Stability — multi-seed spread and bootstrap 95% confidence intervals.",
        ],
    )

    _add_bullets(
        prs,
        "Live demo plan (4–5 min)",
        [
            "Open http://localhost:8501 — Try it live tab.",
            "Run Politics, Sports, Tech, Finance quick examples; then paste Science and Culture snippets.",
            "Point out human-readable topic names vs. topic_id and top keywords.",
            "Switch to Comparison for NMI/Purity; BERTopic explorer for 2-D topic map.",
            "Close with Stability tab: seed variance vs. bootstrap CIs.",
        ],
    )

    _add_bullets(
        prs,
        "Conclusion",
        [
            "Two complementary views of the same Romanian news corpus.",
            "LDA: strong baseline with explicit K and fast iteration.",
            "BERTopic: richer semantic structure at higher compute cost.",
            "Reproducible CLI pipelines + GUI for exploration and presentation.",
            "Questions?",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved -> {OUT}")
    return OUT


if __name__ == "__main__":
    main()
